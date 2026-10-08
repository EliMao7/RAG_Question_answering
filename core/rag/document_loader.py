"""文档解析：把 PDF / DOCX / Markdown / TXT 转成纯文本。

对应技术设计文档 2.1 节「知识库构建流程」的第 ① 步。

支持格式（按需扩展）：

===========  ==================  ==========================================
格式          解析库               说明
===========  ==================  ==========================================
``.pdf``      ``pymupdf``         按页读取；1.24+ 用 ``import pymupdf``，
                                  旧的 ``import fitz`` 已废弃
``.docx``     ``python-docx``     读段落 + 表格单元格
``.md/.txt``  内置                直接读取，统一换行符
===========  ==================  ==========================================

清洗规则：去除页码/页眉类孤立数字行、压缩多余空行、统一换行符。
**保留代码缩进**（缩进即语义，不能压缩）。
"""
from __future__ import annotations

import re
from pathlib import Path

from utils.logger import get_logger

logger = get_logger(__name__)

# 支持的扩展名
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt"}

# 匹配「纯数字行」（常见于页脚页码），清洗时剔除
PAGE_NUMBER_LINE = re.compile(r"^\s*\d{1,4}\s*$")

# 连续 3 个以上换行压缩为 2 个（保留段落感）
MULTI_NEWLINE = re.compile(r"\n{3,}")


def _clean(text: str) -> str:
    """清洗文本：去页码行、压缩空行、统一换行符，但保留代码缩进。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line for line in text.split("\n") if not PAGE_NUMBER_LINE.match(line)]
    text = "\n".join(lines)
    text = MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


def _load_pdf(path: Path) -> str:
    """解析 PDF，按页拼接文本。

    对纯图片页（扫描件）提取不到文字，会跳过——若整份 PDF 都提取不到文本，
    说明是扫描版，需要先做 OCR（技术文档 3.3.1 已注明）。
    """
    import pymupdf  # 1.24+ 的正式包名；旧代码里的 import fitz 已废弃

    doc = pymupdf.open(str(path))
    pages: list[str] = []
    empty_pages = 0

    for page_no, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if not text:
            empty_pages += 1
            continue
        # 插入页标记：便于人工核对分块结果，也让模型知道内容出处
        pages.append(f"[第{page_no}页]\n{text}")

    total_pages = doc.page_count
    doc.close()

    if not pages:
        raise ValueError(f"PDF 未提取到任何文本（可能是扫描版，需先 OCR）：{path.name}")

    logger.info("PDF 解析完成：%s，%s 页（空页 %s 页）", path.name, total_pages, empty_pages)
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    """解析 DOCX，读取段落与表格。"""
    import docx

    document = docx.Document(str(path))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    # 表格内容常包含关键信息（如课程大纲的课时安排），不能漏
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    if not blocks:
        raise ValueError(f"DOCX 未提取到任何文本：{path.name}")

    logger.info("DOCX 解析完成：%s，%s 个文本块（含 %s 个表格）",
                path.name, len(blocks), len(document.tables))
    return "\n\n".join(blocks)


def _load_plain(path: Path) -> str:
    """解析 Markdown / 纯文本。"""
    text = path.read_text(encoding="utf-8", errors="replace")
    logger.info("文本解析完成：%s，%s 字符", path.name, len(text))
    return text


def load_document(path: str | Path) -> str:
    """解析单个文档，返回清洗后的纯文本。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"不支持的文件格式：{suffix}。当前支持 {sorted(SUPPORTED_SUFFIXES)}"
        )

    if suffix == ".pdf":
        return _clean(_load_pdf(path))
    if suffix == ".docx":
        return _clean(_load_docx(path))
    return _clean(_load_plain(path))


def load_directory(
    directory: str | Path,
    recursive: bool = True,
) -> list[tuple[Path, str]]:
    """解析目录下所有受支持的文档。

    :return: ``[(文件路径, 纯文本), ...]``，按文件名排序保证多次运行结果一致
    """
    directory = Path(directory)
    if not directory.is_dir():
        raise NotADirectoryError(f"目录不存在：{directory}")

    pattern = "**/*" if recursive else "*"
    files = sorted(
        p for p in directory.glob(pattern)
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )

    if not files:
        raise FileNotFoundError(
            f"目录中没有可解析的文档：{directory}（支持 {sorted(SUPPORTED_SUFFIXES)}）"
        )

    documents: list[tuple[Path, str]] = []
    for file_path in files:
        try:
            documents.append((file_path, load_document(file_path)))
        except Exception as exc:  # noqa: BLE001 - 单个文件失败不应中断整批建库
            logger.error("解析失败，已跳过：%s | %s", file_path.name, exc)

    logger.info("目录解析完成：%s，成功 %s/%s 个文件",
                directory, len(documents), len(files))
    return documents
